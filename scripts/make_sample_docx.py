"""One-off: build a more realistic sample .docx for the demo-video pilot.

api/test-files/test.docx is a minimal pytest fixture (one line, no real
content) - fine for tests, too sparse to look like a real document on
camera. This writes a standalone sample into scratch-demo-videos/ without
touching the tracked test fixture.
"""

import docx

d = docx.Document()
d.add_heading("Q3 Marketing Report", level=1)
d.add_paragraph(
    "This report summarizes campaign performance for Q3 2026, covering "
    "organic search, paid acquisition, and referral traffic across all "
    "regions."
)
d.add_heading("Key Highlights", level=2)
for item in [
    "Organic traffic grew 18% quarter-over-quarter.",
    "Cost per acquisition dropped from $4.20 to $3.65.",
    "Referral partnerships contributed 12% of new signups.",
]:
    d.add_paragraph(item, style="List Bullet")
d.add_paragraph(
    "Full breakdown by channel is available in the attached spreadsheet. "
    "Next quarter's plan focuses on doubling down on referral growth."
)
d.save("scratch-demo-videos/sample.docx")
print("wrote scratch-demo-videos/sample.docx")
